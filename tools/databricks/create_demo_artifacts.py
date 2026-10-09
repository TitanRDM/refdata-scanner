# Databricks notebook source
# MAGIC %md
# MAGIC # Create refdata-scanner demo artifacts
# MAGIC
# MAGIC **For testing and demos only.** Fills a workspace with realistic reference data for the scanner to find:
# MAGIC
# MAGIC | Where | What gets created |
# MAGIC |---|---|
# MAGIC | Schema `<catalog>.<schema>` | A volume `uploads` with CSV and Excel files; lookup tables (some loaded from those files, so lineage records it); a large fact table that should *not* be flagged; views with `CASE` mappings and `IN` lists |
# MAGIC | Workspace folder | Python and SQL notebooks with hardcoded dicts, lists, `when()` chains, `createDataFrame` literals, `CASE` and `INSERT ... VALUES`, including copies of the same mapping that disagree; a plain `.py` file; CSV and Excel files |
# MAGIC | Optional | A job that references one notebook (never run), a saved SQL query with a long `IN` list, a CSV in DBFS `/FileStore` |
# MAGIC
# MAGIC Every object is listed in `_demo_manifest.json` in the workspace folder. Set **action** to `cleanup` and run again
# MAGIC to delete everything this notebook created.
# MAGIC
# MAGIC **Needs:** a cluster or serverless compute with Unity Catalog, and permission to create a schema (or use an
# MAGIC existing one) in the chosen catalog. Lineage and query history in system tables can take a few hours to appear.
# MAGIC
# MAGIC Unlike the scanner, this notebook **writes** to your workspace. Run it only where test objects are welcome.

# COMMAND ----------

dbutils.widgets.dropdown("action", "create", ["create", "cleanup"], "1. Action")
dbutils.widgets.text("catalog", "", "2. Catalog (blank = current)")
dbutils.widgets.text("schema", "refdata_scanner_demo", "3. Schema to create")
dbutils.widgets.text("workspace_folder", "", "4. Workspace folder (blank = home/refdata-scanner-demo)")
dbutils.widgets.dropdown("create_job", "yes", ["yes", "no"], "5. Create a demo job (never run)")
dbutils.widgets.dropdown("create_query", "yes", ["yes", "no"], "6. Create a saved SQL query")
dbutils.widgets.dropdown("use_dbfs", "yes", ["yes", "no"], "7. Put a CSV in DBFS /FileStore")

# COMMAND ----------

import io
import json
import os
import re

from databricks.sdk import WorkspaceClient
from databricks.sdk.service.workspace import ImportFormat, Language

w = WorkspaceClient()
me = w.current_user.me().user_name

ACTION = dbutils.widgets.get("action")
CAT = dbutils.widgets.get("catalog").strip() or spark.sql("SELECT current_catalog()").first()[0]
SCH = dbutils.widgets.get("schema").strip() or "refdata_scanner_demo"
for _name in (CAT, SCH):
    if not re.match(r"^[A-Za-z0-9_\-]+$", _name):
        raise ValueError(f"Unexpected characters in {_name!r}")
WS = dbutils.widgets.get("workspace_folder").strip().rstrip("/") or f"/Users/{me}/refdata-scanner-demo"
WS = re.sub(r"^/Workspace(?=/)", "", WS)
VOL = f"/Volumes/{CAT}/{SCH}/uploads"
MANIFEST = f"/Workspace{WS}/_demo_manifest.json"
FQ = f"`{CAT}`.`{SCH}`"
print(f"action={ACTION}  schema={CAT}.{SCH}  workspace folder={WS}")

# COMMAND ----------

# MAGIC %md
# MAGIC ## Content

# COMMAND ----------

STATES = [("NSW", "New South Wales"), ("QLD", "Queensland"), ("VIC", "Victoria"), ("SA", "South Australia"),
          ("WA", "Western Australia"), ("TAS", "Tasmania"), ("NT", "Northern Territory"),
          ("ACT", "Australian Capital Territory")]

CSV_FILES = {
    "cost_centres.csv": "cost_centre,name,division,owner\n" + "\n".join(
        f"CC{c},{n},{d},{o}" for c, n, d, o in [
            (100, "Executive", "Corporate", "CEO"), (110, "Finance", "Corporate", "CFO"),
            (120, "People and Culture", "Corporate", "CHRO"), (130, "Legal", "Corporate", "GC"),
            (200, "Retail Sales", "Sales", "Head of Retail"), (210, "Wholesale Sales", "Sales", "Head of Wholesale"),
            (220, "Online Sales", "Sales", "Head of Digital"), (300, "Warehouse North", "Operations", "Ops Manager"),
            (310, "Warehouse South", "Operations", "Ops Manager"), (320, "Fleet", "Operations", "Fleet Manager"),
            (400, "IT", "Technology", "CIO"), (410, "Data", "Technology", "CDO")]) + "\n",
    "product_categories.csv": "product_code,category,sub_category\n" + "\n".join(
        f"P{100 + i},{cat},{sub}" for i, (cat, sub) in enumerate(
            [("Hardware", "Laptops"), ("Hardware", "Monitors"), ("Hardware", "Accessories"),
             ("Software", "Licences"), ("Software", "Subscriptions"), ("Services", "Support"),
             ("Services", "Installation"), ("Services", "Training")])) + "\n",
    "channel_codes.csv": "channel_code,channel_name\n01,Retail\n02,Wholesale\n03,Online\n04,Marketplace\n05,Phone\n",
    "regions/state_regions.csv": "state,region\nNSW,East\nVIC,East\nQLD,North\nNT,North\nWA,West\nSA,South\nTAS,South\nACT,East\n",
}

XLSX_ROWS = [("gl_account", "description", "report_line"), ("4000", "Product sales", "Revenue"),
             ("4100", "Service revenue", "Revenue"), ("5000", "Cost of goods sold", "Cost of Sales"),
             ("6000", "Salaries and wages", "Operating Expenses"), ("6100", "Rent", "Operating Expenses"),
             ("7000", "Interest income", "Other Income")]

NB_SALES = '''#__HEADER__
#__MAGIC__ %md
#__MAGIC__ # Sales transform
#__MAGIC__ Cleans raw orders and adds descriptive columns.

#__CMD__

from pyspark.sql import functions as F

STATE_MAP = {
__STATES__
}

select_cols = ["order_id", "order_date", "customer_id", "state", "channel_code", "product_code",
               "qty", "unit_price", "discount", "tax", "total", "currency"]

#__CMD__

orders = spark.table("__FQ__.sales_orders").select(*select_cols)
orders = orders.replace(STATE_MAP, subset=["state"])

orders = orders.withColumn(
    "channel",
    F.when(F.col("channel_code") == "01", "Retail")
     .when(F.col("channel_code") == "02", "Wholesale")
     .when(F.col("channel_code") == "03", "Online")
     .when(F.col("channel_code") == "04", "Marketplace")
     .when(F.col("channel_code") == "05", "Phone")
     .otherwise("Unknown"),
)

categories = spark.read.option("header", True).csv("__VOL__/product_categories.csv")
orders = orders.join(categories, "product_code", "left")

#__CMD__

#__MAGIC__ %sql
#__MAGIC__ SELECT *,
#__MAGIC__   CASE WHEN product_code IN ('P100','P101','P102','P103','P104','P105','P106','P107','P108','P109','P110','P111')
#__MAGIC__        THEN 'Discontinued' ELSE 'Active' END AS product_status
#__MAGIC__ FROM __FQ__.sales_orders
'''

NB_FINANCE = '''#__HEADER__
import pandas as pd
from pyspark.sql import functions as F

# Copied from the sales notebook. TODO: keep in sync
STATE_MAP = {
__STATES_DRIFT__
}

COST_CENTRES = ["CC100", "CC110", "CC120", "CC130", "CC200", "CC210", "CC220", "CC300", "CC310", "CC320", "CC400", "CC410"]

spark_conf = {"spark.sql.shuffle.partitions": "64", "spark.sql.adaptive.enabled": "true",
              "spark.databricks.delta.optimizeWrite.enabled": "true", "spark.sql.ansi.enabled": "false",
              "spark.sql.session.timeZone": "Australia/Brisbane"}

#__CMD__

gl = spark.table("__FQ__.sales_orders").filter(F.col("cost_centre").isin(COST_CENTRES))
mapping = pd.read_excel("/Workspace__WS__/finance/gl_account_mapping.xlsx")

#__CMD__

account_groups = spark.createDataFrame(
    [
        ("4000", "Revenue"),
        ("4100", "Revenue"),
        ("5000", "Cost of Sales"),
        ("6000", "Operating Expenses"),
        ("6100", "Operating Expenses"),
        ("7000", "Other Income"),
    ],
    ["account", "account_group"],
)
'''

NB_REGION_SQL = '''--__HEADER__
CREATE OR REPLACE TEMP VIEW sales_by_region AS
SELECT
  CASE state
    WHEN 'NSW' THEN 'East'
    WHEN 'VIC' THEN 'East'
    WHEN 'QLD' THEN 'North'
    WHEN 'NT'  THEN 'North'
    WHEN 'WA'  THEN 'West'
    WHEN 'SA'  THEN 'South'
    WHEN 'TAS' THEN 'South'
    ELSE 'Other'
  END AS sales_region,
  sum(total) AS total
FROM __FQ__.sales_orders
GROUP BY 1

--__CMD__

--__MAGIC__ %python
--__MAGIC__ spark.sql("""
--__MAGIC__   INSERT INTO __FQ__.ref_currency VALUES
--__MAGIC__     ('AUD', 'Australian Dollar'), ('NZD', 'New Zealand Dollar'), ('USD', 'US Dollar'),
--__MAGIC__     ('EUR', 'Euro'), ('GBP', 'Pound Sterling')
--__MAGIC__ """)

--__CMD__

SELECT * FROM read_files('__VOL__/cost_centres.csv', format => 'csv', header => true)
'''

NB_MARKETING = '''#__HEADER__
from pyspark.sql import functions as F

# Same codes as the sales notebook's when() chain, kept separately here
CHANNELS = {"01": "Retail", "02": "Wholesale", "03": "Online", "04": "Marketplace", "05": "Phone"}

SEGMENT_GROUPS = {
    "Business": ["E", "M", "S"],
    "Public": ["G", "EDU"],
    "Consumer": ["C"],
    "Partner": ["R", "D"],
    "Internal": ["X"],
}

COUNTRY_NAMES = {
    "AU": "Australia", "NZ": "New Zealand", "US": "United States", "GB": "United Kingdom", "CA": "Canada",
    "SG": "Singapore", "JP": "Japan", "DE": "Germany", "FR": "France", "IN": "India", "CN": "China",
    "ID": "Indonesia", "MY": "Malaysia", "PH": "Philippines", "TH": "Thailand",
}

#__CMD__

customers = spark.table("__FQ__.sales_orders").withColumn(
    "tier",
    F.when(F.col("tier_code") == "P", "Platinum")
     .when(F.col("tier_code") == "G", "Gold")
     .when(F.col("tier_code") == "S", "Silver")
     .when(F.col("tier_code") == "B", "Bronze")
     .when(F.col("tier_code") == "N", "New")
     .otherwise("Unknown"),
)
campaign_codes = spark.read.csv("/Workspace__WS__/marketing/uploads/campaign_codes.csv", header=True)
'''

PY_MAPPINGS = '''"""Shared lookups. Edit here and redeploy."""

ORDER_STATUS = {
    "N": "New",
    "P": "Picking",
    "S": "Shipped",
    "D": "Delivered",
    "R": "Returned",
    "C": "Cancelled",
}

UOM_TO_EACH = {"EA": 1, "PR": 2, "DZ": 12, "BX": 24, "CS": 48, "PL": 960}
'''

CAMPAIGN_CSV = "campaign_code,campaign_name,channel\nSPR26,Spring sale,Online\nEOFY,End of financial year,All\nBTS,Back to school,Retail\n"


def _states_dict(drift=False):
    rows = []
    for code, name in STATES:
        if drift and code == "WA":
            name = "West Australia"
        if drift and code == "ACT":
            name = "ACT"
        rows.append(f'    "{code}": "{name}",')
    return "\n".join(rows)


def render(text):
    # Notebook markers are kept as placeholders above so this notebook's own cells aren't split on import.
    text = (text.replace("__CMD__", " COMMAND ----------").replace("__MAGIC__", " MAGIC")
                .replace("__HEADER__", " Databricks notebook source"))
    return (text.replace("__STATES_DRIFT__", _states_dict(drift=True)).replace("__STATES__", _states_dict())
            .replace("__FQ__", f"{CAT}.{SCH}").replace("__VOL__", VOL).replace("__WS__", WS))


def xlsx_bytes(rows):
    try:
        import openpyxl
    except ImportError:
        return None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mapping"
    for r in rows:
        ws.append(list(r))
    wb.create_sheet("Notes")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# COMMAND ----------

# MAGIC %md
# MAGIC ## Create or clean up

# COMMAND ----------

def create():
    manifest = {"schema": f"{CAT}.{SCH}", "workspace_folder": WS, "job_id": None, "query_id": None, "dbfs": []}

    # --- schema, volume and files ------------------------------------------------
    spark.sql(f"CREATE SCHEMA IF NOT EXISTS {FQ} COMMENT 'refdata-scanner demo objects; safe to drop'")
    spark.sql(f"CREATE VOLUME IF NOT EXISTS {FQ}.uploads COMMENT 'refdata-scanner demo uploads'")
    for rel, text in CSV_FILES.items():
        path = f"{VOL}/{rel}"
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w") as fh:
            fh.write(text)
    xlsx = xlsx_bytes(XLSX_ROWS)
    if xlsx:
        with open(f"{VOL}/gl_account_mapping.xlsx", "wb") as fh:
            fh.write(xlsx)
    else:
        print("openpyxl not available: Excel files skipped")
    print(f"files written to {VOL}")

    # --- tables -----------------------------------------------------------------------
    spark.sql(f"""CREATE OR REPLACE TABLE {FQ}.lkp_order_status (status_code STRING, status_name STRING,
                  is_open BOOLEAN, sort_order INT) COMMENT 'Order status lookup'""")
    spark.sql(f"""INSERT INTO {FQ}.lkp_order_status VALUES ('N','New',true,1),('P','Picking',true,2),
                  ('S','Shipped',true,3),('D','Delivered',false,4),('R','Returned',false,5),('C','Cancelled',false,6)""")
    # Loaded from files, so lineage records the CSV as the source
    spark.sql(f"""CREATE OR REPLACE TABLE {FQ}.dim_channel AS
                  SELECT * FROM read_files('{VOL}/channel_codes.csv', format => 'csv', header => true)""")
    spark.sql(f"""CREATE OR REPLACE TABLE {FQ}.xref_state_region AS
                  SELECT * FROM read_files('{VOL}/regions/state_regions.csv', format => 'csv', header => true)""")
    spark.sql(f"CREATE OR REPLACE TABLE {FQ}.ref_cost_centre (cost_centre STRING, name STRING, division STRING, owner STRING)")
    spark.sql(f"""COPY INTO {FQ}.ref_cost_centre FROM '{VOL}/cost_centres.csv'
                  FILEFORMAT = CSV FORMAT_OPTIONS ('header' = 'true')""")
    spark.sql(f"CREATE OR REPLACE TABLE {FQ}.ref_currency (currency_code STRING, currency_name STRING)")
    # A wide fact table: should NOT be reported as a lookup table
    cols = ", ".join([f"CAST(id % 97 AS DOUBLE) AS measure_{i}" for i in range(20)])
    spark.sql(f"""CREATE OR REPLACE TABLE {FQ}.sales_orders AS
                  SELECT id AS order_id, date_add('2026-01-01', CAST(id % 270 AS INT)) AS order_date,
                         element_at(array('NSW','QLD','VIC','SA','WA','TAS','NT','ACT'), CAST(id % 8 AS INT) + 1) AS state,
                         lpad(CAST(id % 5 + 1 AS STRING), 2, '0') AS channel_code,
                         concat('P', CAST(100 + id % 8 AS STRING)) AS product_code,
                         concat('CC', CAST(100 + (id % 4) * 10 AS STRING)) AS cost_centre,
                         element_at(array('P','G','S','B','N'), CAST(id % 5 AS INT) + 1) AS tier_code,
                         {cols}
                  FROM range(5000)""")

    # --- views with hardcoded mappings --------------------------------------------------
    spark.sql(f"""CREATE OR REPLACE VIEW {FQ}.v_sales_by_region AS
                  SELECT CASE state WHEN 'NSW' THEN 'East' WHEN 'VIC' THEN 'East' WHEN 'QLD' THEN 'North'
                                    WHEN 'NT' THEN 'North' WHEN 'WA' THEN 'West' WHEN 'SA' THEN 'South'
                                    WHEN 'TAS' THEN 'South' ELSE 'Other' END AS sales_region,
                         sum(measure_0) AS total
                  FROM {FQ}.sales_orders GROUP BY 1""")
    spark.sql(f"""CREATE OR REPLACE VIEW {FQ}.v_active_products AS
                  SELECT * FROM {FQ}.sales_orders
                  WHERE product_code NOT IN ('P100','P101','P102','P103','P104','P105','P106','P107','P108',
                                             'P109','P110','P111')""")
    spark.sql(f"""CREATE OR REPLACE VIEW {FQ}.v_order_status AS
                  SELECT order_id, CASE tier_code WHEN 'P' THEN 'Platinum' WHEN 'G' THEN 'Gold' WHEN 'S' THEN 'Silver'
                                                  WHEN 'B' THEN 'Bronze' WHEN 'N' THEN 'New' ELSE 'Unknown' END AS tier
                  FROM {FQ}.sales_orders""")
    print(f"tables and views created in {CAT}.{SCH}")

    # --- workspace notebooks and files --------------------------------------------------
    for sub in ("sales", "finance", "marketing/uploads", "shared_utils"):
        w.workspace.mkdirs(f"{WS}/{sub}")
    notebooks = [("sales/sales_transform", NB_SALES, Language.PYTHON),
                 ("finance/finance_report", NB_FINANCE, Language.PYTHON),
                 ("finance/region_rollup", NB_REGION_SQL, Language.SQL),
                 ("marketing/campaign_segments", NB_MARKETING, Language.PYTHON)]
    for rel, text, lang in notebooks:
        w.workspace.upload(f"{WS}/{rel}", render(text).encode(), format=ImportFormat.SOURCE, language=lang,
                           overwrite=True)
    w.workspace.upload(f"{WS}/shared_utils/mappings.py", PY_MAPPINGS.encode(), format=ImportFormat.AUTO, overwrite=True)
    w.workspace.upload(f"{WS}/marketing/uploads/campaign_codes.csv", CAMPAIGN_CSV.encode(), format=ImportFormat.AUTO,
                       overwrite=True)
    if xlsx:
        w.workspace.upload(f"{WS}/finance/gl_account_mapping.xlsx", xlsx, format=ImportFormat.AUTO, overwrite=True)
    print(f"notebooks and files created in {WS}")

    # --- optional extras ------------------------------------------------------------------
    if dbutils.widgets.get("create_job") == "yes":
        try:
            from databricks.sdk.service.jobs import NotebookTask, Task

            job = w.jobs.create(name="refdata-scanner demo: finance report (do not run)",
                                tags={"refdata_scanner_demo": "true"},
                                tasks=[Task(task_key="finance_report",
                                            notebook_task=NotebookTask(notebook_path=f"{WS}/finance/finance_report"))])
            manifest["job_id"] = job.job_id
            print(f"job created: {job.job_id} (it is never run)")
        except Exception as exc:
            print(f"job not created: {exc}")

    if dbutils.widgets.get("create_query") == "yes":
        try:
            from databricks.sdk.service.sql import CreateQueryRequestQuery

            warehouse = next(iter(w.warehouses.list()), None)
            if warehouse is None:
                raise RuntimeError("no SQL warehouse available")
            q = w.queries.create(query=CreateQueryRequestQuery(
                display_name="refdata-scanner demo: channel report", warehouse_id=warehouse.id, parent_path=WS,
                query_text=f"SELECT * FROM {CAT}.{SCH}.sales_orders WHERE cost_centre IN ('CC100','CC110','CC120',"
                           "'CC130','CC200','CC210','CC220','CC300','CC310','CC320','CC400','CC410')"))
            manifest["query_id"] = q.id
            print(f"saved query created: {q.id}")
        except Exception as exc:
            print(f"saved query not created: {exc}")

    if dbutils.widgets.get("use_dbfs") == "yes":
        try:
            path = "dbfs:/FileStore/tables/refdata_scanner_demo/uom_codes.csv"
            dbutils.fs.put(path, "uom,description,each\nEA,Each,1\nDZ,Dozen,12\nBX,Box,24\nPL,Pallet,960\n", True)
            manifest["dbfs"].append(path)
            print(f"DBFS file written: {path}")
        except Exception as exc:
            print(f"DBFS file not written (DBFS may be disabled): {exc}")

    with open(MANIFEST, "w") as fh:
        json.dump(manifest, fh, indent=2)
    print(f"\nDone. Manifest: {MANIFEST}")
    print(f"Now run 02_full_scan with catalog={CAT}, schemas={SCH}, workspace folders={WS}")


def cleanup():
    manifest = {}
    if os.path.exists(MANIFEST):
        with open(MANIFEST) as fh:
            manifest = json.load(fh)
    if manifest.get("job_id"):
        try:
            w.jobs.delete(manifest["job_id"])
            print(f"job {manifest['job_id']} deleted")
        except Exception as exc:
            print(f"job not deleted: {exc}")
    if manifest.get("query_id"):
        try:
            w.queries.delete(manifest["query_id"])
            print(f"query {manifest['query_id']} deleted")
        except Exception as exc:
            print(f"query not deleted: {exc}")
    for path in manifest.get("dbfs", []):
        try:
            dbutils.fs.rm(path)
        except Exception:
            pass
    spark.sql(f"DROP SCHEMA IF EXISTS {FQ} CASCADE")
    print(f"schema {CAT}.{SCH} dropped")
    try:
        w.workspace.delete(WS, recursive=True)
        print(f"workspace folder {WS} deleted")
    except Exception as exc:
        print(f"workspace folder not deleted: {exc}")


create() if ACTION == "create" else cleanup()
