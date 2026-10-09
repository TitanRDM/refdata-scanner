"""Optional, opt-in peek inside CSV/Excel files: header row, row estimate, sheet count.

Only the first ``max_inspect_bytes`` of a CSV are read. Excel files are only
opened when they are smaller than that limit. Cell values are never stored,
only column headers.
"""
from __future__ import annotations

import csv
import io
from typing import Optional

from .models import FileAsset


def inspect_file(f: FileAsset, data: Optional[bytes]) -> None:
    if not data:
        return
    try:
        if f.extension in (".csv", ".tsv"):
            _inspect_csv(f, data)
        elif f.extension in (".xlsx", ".xlsm"):
            _inspect_xlsx(f, data)
    except Exception:  # a broken file should never stop the scan
        pass


def _inspect_csv(f: FileAsset, data: bytes) -> None:
    text = data.decode("utf-8-sig", errors="replace")
    lines = text.splitlines()
    if not lines:
        return
    delimiter = "\t" if f.extension == ".tsv" else ","
    try:
        delimiter = csv.Sniffer().sniff(lines[0][:2000], delimiters=",;\t|").delimiter
    except csv.Error:
        pass
    f.header = [h.strip() for h in next(csv.reader([lines[0]], delimiter=delimiter))][:50]
    complete = len(lines) - (0 if text.endswith("\n") else 1)
    if f.size_bytes and len(data) < f.size_bytes and complete > 1:
        f.row_estimate = int((complete - 1) * f.size_bytes / len(data))
    else:
        f.row_estimate = max(len(lines) - 1, 0)


def _inspect_xlsx(f: FileAsset, data: bytes) -> None:
    if f.size_bytes and len(data) < f.size_bytes:
        return  # truncated zip cannot be opened
    try:
        import openpyxl  # optional dependency
    except ImportError:
        return
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    f.sheet_count = len(wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    first = next(ws.iter_rows(max_row=1, values_only=True), None)
    if first:
        f.header = [str(v).strip() for v in first if v is not None][:50]
    if ws.max_row:
        f.row_estimate = max(ws.max_row - 1, 0)
    wb.close()
